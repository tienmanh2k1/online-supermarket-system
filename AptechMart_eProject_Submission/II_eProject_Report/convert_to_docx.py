#!/usr/bin/env python3
"""
Convert markdown files to DOCX format with proper formatting.
Handles tables, code blocks, headings, and basic styling.
"""

import os
import re
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def create_docx(md_files, output_file):
    """Convert multiple markdown files to a single DOCX document."""
    doc = Document()

    # Set default styles
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)

    for md_file in md_files:
        with open(md_file, 'r', encoding='utf-8') as f:
            content = f.read()

        process_markdown_to_docx(doc, content)

    doc.save(output_file)
    print(f"[OK] Document saved: {output_file}")

def process_markdown_to_docx(doc, content):
    """Process markdown content and add to document."""
    lines = content.split('\n')
    in_code_block = False
    code_block_content = []
    code_block_lang = ""

    i = 0
    while i < len(lines):
        line = lines[i]

        # Code block handling
        if line.strip().startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_block_lang = line.strip()[3:]
                code_block_content = []
            else:
                # End of code block - add to doc
                add_code_block(doc, '\n'.join(code_block_content), code_block_lang)
                in_code_block = False
            i += 1
            continue
        elif in_code_block:
            code_block_content.append(line)
            i += 1
            continue

        # Table handling
        if '|' in line and line.strip().startswith('|'):
            # Collect all table lines
            table_lines = []
            while i < len(lines) and '|' in lines[i] and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue

        # Image handling
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line.strip())
        if img_match:
            caption = img_match.group(1)
            img_rel = img_match.group(2)
            base_dir = os.path.dirname(os.path.abspath(__file__))
            img_full = os.path.normpath(os.path.join(base_dir, img_rel))
            if os.path.exists(img_full):
                try:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    doc.add_picture(img_full, width=Inches(6.0))
                    if caption:
                        cp = doc.add_paragraph()
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        crun = cp.add_run(caption)
                        crun.italic = True
                        crun.font.size = Pt(9.5)
                        crun.font.color.rgb = RGBColor(71, 85, 105)
                except Exception as e:
                    print(f"[WARN] Failed to insert image {img_full}: {e}")
            else:
                print(f"[WARN] Image not found: {img_full}")
            i += 1
            continue

        # Heading handling
        if line.startswith('# '):
            add_heading(doc, line[2:], 1)
        elif line.startswith('## '):
            add_heading(doc, line[3:], 2)
        elif line.startswith('### '):
            add_heading(doc, line[4:], 3)
        elif line.startswith('#### '):
            add_heading(doc, line[5:], 4)
        # Horizontal rule
        elif line.strip() in ['---', '***', '___']:
            add_horizontal_rule(doc)
        # Empty line
        elif line.strip() == '':
            doc.add_paragraph()
        # Regular paragraph
        else:
            add_paragraph(doc, line)

        i += 1

def add_heading(doc, text, level):
    """Add a heading with appropriate style."""
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_paragraph(doc, text):
    """Add a paragraph with inline formatting."""
    # Handle bold (**text**)
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)

    para = doc.add_paragraph()
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = para.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = para.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
        else:
            # Handle Vietnamese text
            run = para.add_run(part)

def add_code_block(doc, code, language):
    """Add a code block with monospace font."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.3)

    run = para.add_run(code)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)

    # Add light gray background (simulated with border)
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '4')
    left.set(qn('w:space'), '4')
    left.set(qn('w:color'), 'AAAAAA')
    pBdr.append(left)
    pPr.append(pBdr)

    # Light background color
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F5F5F5')
    pPr.append(shd)

def add_table(doc, table_lines):
    """Add a table from markdown table lines."""
    if len(table_lines) < 2:
        return

    # Parse header
    header = [cell.strip() for cell in table_lines[0].split('|')[1:-1]]

    # Count columns
    num_cols = len(header)

    # Skip separator line
    start_row = 1
    if len(table_lines) > 1 and re.match(r'\|[\s\-:|]+\|', table_lines[1]):
        start_row = 2

    # Create table
    table = doc.add_table(rows=0, cols=num_cols)
    table.style = 'Table Grid'

    # Add header row
    header_row = table.add_row()
    for i, cell_text in enumerate(header):
        cell = header_row.cells[i]
        cell.text = cell_text
        # Bold header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    # Add data rows
    for row_idx in range(start_row, len(table_lines)):
        row_data = [cell.strip() for cell in table_lines[row_idx].split('|')[1:-1]]
        if len(row_data) == num_cols:
            row = table.add_row()
            for i, cell_text in enumerate(row_data):
                row.cells[i].text = cell_text

    doc.add_paragraph()

def add_horizontal_rule(doc):
    """Add a horizontal line."""
    para = doc.add_paragraph()
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '888888')
    pBdr.append(bottom)
    pPr.append(pBdr)

if __name__ == '__main__':
    # Convert both main report and supplements
    md_files = [
        'eProject_Report.md',
        'CHAPTER_SUPPLEMENTS.md'
    ]
    output = 'eProject_Report_Final.docx'

    create_docx(md_files, output)
